"""Reads submitted concept files for the in-app viewer and the AI analysis.

The files come from customers, so every read is bounded: archive size, page count and extracted text.
"""
import io
import re
import zipfile
from dataclasses import dataclass, field

from docx import Document
from docx.table import Table
from lxml.etree import LxmlError
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader
from pypdf.errors import PyPdfError

MAX_UNZIPPED_MB = 150  # A Word or PowerPoint file is a zip; refuse archives that expand beyond this.
MAX_PDF_PAGES = 150
MAX_TEXT_CHARS = 200_000
MAX_TABLE_ROWS = 200

READABLE_EXTENSIONS = {"pdf": "pdf", "docx": "word", "pptx": "slides"}


class ConceptFileError(Exception):
    """Safe, user-facing reason a concept file cannot be read."""


@dataclass
class ConceptContent:
    """Extracted file content.

    word blocks: {"type": "heading", "level", "text"}, {"type": "paragraph", "text", "list"}, {"type": "table", "rows"}
    slides blocks: {"number", "title", "lines", "notes"}
    pdf blocks: {"number", "text"}
    """

    kind: str
    blocks: list = field(default_factory=list)
    truncated: bool = False
    _chars: int = 0

    def add_text(self, text):
        """Count extracted text; returns False once the limit is reached."""
        self._chars += len(text)
        if self._chars > MAX_TEXT_CHARS:
            self.truncated = True
            return False
        return True

    def as_text(self):
        parts = []
        for block in self.blocks:
            if self.kind == "word":
                if block["type"] == "heading":
                    parts.append("#" * min(block["level"], 6) + " " + block["text"])
                elif block["type"] == "table":
                    parts.append("\n".join(" | ".join(row) for row in block["rows"]))
                else:
                    parts.append(("- " if block["list"] else "") + block["text"])
            elif self.kind == "slides":
                lines = [f"--- Slide {block['number']}: {block['title'] or 'Untitled'} ---", *block["lines"]]
                if block["notes"]:
                    lines.append("Speaker notes: " + block["notes"])
                parts.append("\n".join(lines))
            else:
                parts.append(f"--- Page {block['number']} ---\n{block['text']}")
        return "\n\n".join(parts).strip()


def _clean(text):
    return re.sub(r"[ \t ]+", " ", text or "").strip()


def _read_bytes(field_file):
    try:
        with field_file.storage.open(field_file.name, "rb") as handle:
            return io.BytesIO(handle.read())
    except FileNotFoundError as exc:
        raise ConceptFileError("The file is missing from the server.") from exc


def _check_archive(data):
    try:
        with zipfile.ZipFile(data) as archive:
            expanded = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise ConceptFileError("This file is damaged or is not a valid Word or PowerPoint document.") from exc
    if expanded > MAX_UNZIPPED_MB * 1024 * 1024:
        raise ConceptFileError("This document is too large to preview. Download it to open it.")
    data.seek(0)


def _read_word(data):
    content = ConceptContent("word")
    document = Document(data)
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            rows = []
            for row in item.rows[:MAX_TABLE_ROWS]:
                cells = []
                for cell in row.cells:
                    text = _clean(cell.text)
                    # Merged cells are reported once per grid column; show them once.
                    if not cells or cells[-1] != text:
                        cells.append(text)
                rows.append(cells)
            if rows and content.add_text(" ".join(" ".join(row) for row in rows)):
                content.blocks.append({"type": "table", "rows": rows})
            continue
        text = _clean(item.text)
        if not text:
            continue
        if not content.add_text(text):
            break
        style = (item.style.name if item.style is not None else "") or ""
        level = re.match(r"Heading (\d)", style)
        if level or style == "Title":
            content.blocks.append({"type": "heading", "level": int(level.group(1)) if level else 1, "text": text})
        else:
            numbered = item._p.pPr is not None and item._p.pPr.numPr is not None
            content.blocks.append({"type": "paragraph", "text": text, "list": numbered or "List" in style})
    return content


def _shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shapes(shape.shapes)
        else:
            yield shape


def _read_slides(data):
    content = ConceptContent("slides")
    presentation = Presentation(data)
    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = _clean(title_shape.text) if title_shape is not None else ""
        lines = []
        for shape in _shapes(slide.shapes):
            if shape == title_shape:
                continue
            if shape.has_text_frame:
                lines += ["  " * paragraph.level + _clean(paragraph.text)
                          for paragraph in shape.text_frame.paragraphs if _clean(paragraph.text)]
            elif getattr(shape, "has_table", False) and shape.has_table:
                lines += [" | ".join(_clean(cell.text) for cell in row.cells) for row in shape.table.rows]
        notes = _clean(slide.notes_slide.notes_text_frame.text) if slide.has_notes_slide else ""
        if not content.add_text(title + "".join(lines) + notes):
            break
        content.blocks.append({"number": number, "title": title, "lines": lines, "notes": notes})
    return content


def _read_pdf(data):
    content = ConceptContent("pdf")
    reader = PdfReader(data)
    if reader.is_encrypted and not reader.decrypt(""):
        raise ConceptFileError("This PDF is password-protected, so its text cannot be read.")
    if len(reader.pages) > MAX_PDF_PAGES:
        content.truncated = True
    for number, page in enumerate(reader.pages[:MAX_PDF_PAGES], start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            continue
        if not content.add_text(text):
            break
        content.blocks.append({"number": number, "text": text})
    return content


def read_concept_file(field_file):
    """Return the readable content of a PDF, .docx or .pptx concept file."""
    if not field_file:
        raise ConceptFileError("The customer did not attach a file.")
    kind = READABLE_EXTENSIONS.get(field_file.name.rsplit(".", 1)[-1].lower())
    if kind is None:
        raise ConceptFileError("Older .doc and .ppt files cannot be read in the app. Download the file to open it.")
    data = _read_bytes(field_file)
    try:
        if kind == "pdf":
            return _read_pdf(data)
        _check_archive(data)
        return _read_word(data) if kind == "word" else _read_slides(data)
    except ConceptFileError:
        raise
    except (PyPdfError, LxmlError, KeyError, ValueError, TypeError, AttributeError, IndexError, zipfile.BadZipFile, OSError) as exc:
        # Library errors for malformed files vary widely; none of their details are shown.
        raise ConceptFileError("This file could not be read. It may be damaged; download it to open it.") from exc
